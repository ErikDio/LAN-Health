import openpyxl
from openpyxl.chart import BarChart, Reference, LineChart
import openpyxl.workbook
import variaveis
import log

from copy import deepcopy

class Grafico:

    arquivo_planilha = None
    
    def setup(self, lcl):
        self.arquivo_planilha = lcl

    def gerar_grafico_barra(titulo_y, dados, categorias):
        grafico = BarChart()
        grafico.type = "col"
        grafico.style = 13
        grafico.grouping = "stacked"
        grafico.overlap = 100
        grafico.y_axis.title = titulo_y
        grafico.x_axis.title = "Horário"
        grafico.width = 30
        grafico.height = 12
        grafico.y_axis.scaling.max = 255
        grafico.y_axis.scaling.min = 0
        grafico.add_data(dados,titles_from_data=True)
        grafico.set_categories(categorias)
        return grafico
    
    def gerar_grafico_linha(titulo, titulo_y, dados, categorias):
        grafico = LineChart()
        grafico.title = titulo
        grafico.style = 13
        grafico.x_axis.title = "Horário"
        grafico.y_axis.title = titulo_y
        grafico.width = 30
        grafico.height = 12
        grafico.y_axis.scaling.min = 0
        grafico.add_data(dados,titles_from_data=True)
        grafico.set_categories(categorias)
        return grafico
    
    def gerar_graficos(self):
        wb = openpyxl.load_workbook(self.arquivo_planilha)        
        wbgrafico = wb["Grafico"]
        wbdados = wb["Dados"]
        
        dados_uso_hoje = Reference(wbdados, min_col=2, max_col=3, min_row=1, max_row=25)
        dados_uso_ontem = Reference(wbdados, min_col=4, max_col=5, min_row=1, max_row=25)
        dados_conflito = Reference(wbdados, min_col=6, max_col=6,min_row=1, max_row=25)
        dados_estimativa = Reference(wbdados, min_col=7, max_col=7,min_row=1, max_row=25)
        dados_estimativa_ontem = Reference(wbdados, min_col=8, max_col=8,min_row=1, max_row=25)

        categorias = Reference(wbdados, min_col=1, max_col=1, min_row=2, max_row=25)
        
        uso_grafico = Grafico.gerar_grafico_barra(titulo_y="Dispositivos Online (Hoje)", dados=dados_uso_hoje, categorias=categorias)
        uso_grafico_ontem = Grafico.gerar_grafico_barra(titulo_y="Dispositivos Online (Ontem)", dados=dados_uso_ontem, categorias=categorias)
        conflito_grafico = Grafico.gerar_grafico_linha(titulo="IPs Conflituosos Ativos", titulo_y="Conflitos", dados=dados_conflito, categorias=categorias)
        estimativa_grafico = Grafico.gerar_grafico_linha(titulo="Velocidade de download (Hoje)", titulo_y="MB/s", dados=dados_estimativa, categorias=categorias)
        estimativa_grafico_ontem = Grafico.gerar_grafico_linha(titulo="Velocidade de download (Ontem)", titulo_y="MB/s", dados=dados_estimativa_ontem, categorias=categorias)

        wbgrafico.add_chart(uso_grafico, f"A{1}")
        wbgrafico.add_chart(uso_grafico_ontem, f"A{1+26}")
        wbgrafico.add_chart(conflito_grafico, f"A{1+26*2}")
        wbgrafico.add_chart(estimativa_grafico, f"A{1+26*3}")
        wbgrafico.add_chart(estimativa_grafico_ontem, f"A{1+26*4}")

        wb.save(self.arquivo_planilha)